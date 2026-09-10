"""SEFAZ-GO da planilha até a entrega, sem portal.

O adapter ter os desfechos certos não basta: o que a operação faz é clicar
na automação e esperar o Excel e o ZIP saírem como saem da Receita. Entre
uma coisa e outra há ingestão, fila, gravação de certidão, relatório e
pacote — e cada uma delas conhece o órgão por um caminho diferente
(`ABA_PARA_ORGAO`, `cfg.orgaos`, `NOMES_DE_ORGAO`).

Estes testes percorrem a cadeia com o adapter simulado. Um órgão novo que
esqueça de um desses pontos passa nos testes de unidade e falha aqui.
"""
from __future__ import annotations

import io
import zipfile

import pytest

from cnd.core import fila, tempo
from cnd.core.modelos import Desfecho, ResultadoTentativa
from cnd.infra.config import carregar, nome_do_orgao
from cnd.infra.db import conectar, criar_schema
from cnd.ingestao.planilha import importar
from cnd.web.relatorio import Recorte, gerar, zipar_pdfs

ORGAO = "SEFAZ_GO"


@pytest.fixture
def carteira(tmp_path):
    """Uma planilha com a aba GO, como a que a operação envia."""
    from openpyxl import Workbook

    livro = Workbook()
    aba = livro.active
    aba.title = "GO"
    aba.append(["Empresa", "CNPJ"])
    for documento, nome in [
        # Dígitos verificadores válidos: a ingestão rejeita o resto, e um
        # CNPJ inventado sem conta certa faria o teste medir a rejeição.
        ("11.222.333/0001-81", "ALFA LTDA"),
        ("22.333.444/0001-81", "BETA SA"),
        ("33.444.555/0001-81", "GAMA ME"),
        ("44.555.666/0001-81", "DELTA EIRELI"),
    ]:
        aba.append([nome, documento])
    caminho = tmp_path / "CARTEIRA.xlsx"
    livro.save(caminho)
    return caminho


@pytest.fixture
def processado(tmp_path, carteira):
    """Importa a aba GO e deixa o robô processar tudo, sem tocar no portal."""
    conn = conectar(tmp_path / "cnd.db")
    criar_schema(conn)
    importar(conn, carteira, "Importacao GO", ["GO"], ORGAO)
    conn.commit()

    pdfs = tmp_path / "pdfs"
    pdfs.mkdir()
    desfechos = [Desfecho.NEGATIVA, Desfecho.POSITIVA,
                 Desfecho.CPEN, Desfecho.NEGATIVA]
    for desfecho in desfechos:
        job = fila.reivindicar(conn, ORGAO)
        assert job is not None, "a ingestão não colocou o item na fila do órgão"
        tentativa = fila.abrir_tentativa(conn, job, 1)
        caminho = None
        if desfecho in (Desfecho.NEGATIVA, Desfecho.CPEN):
            caminho = pdfs / f"{job.doc.documento}.pdf"
            caminho.write_bytes(b"%PDF-1.4 certidao")
        resultado = ResultadoTentativa(desfecho, caminho_pdf=caminho,
                                       codigo_controle="5.555.451",
                                       mensagem_portal="ok")
        fila.fechar_tentativa(conn, tentativa, resultado)
        fila.concluir(conn, job, resultado)
    conn.commit()
    yield conn
    conn.close()


class TestDaPlanilhaAteAFila:
    def test_a_aba_go_vira_fila_do_orgao(self, tmp_path, carteira):
        conn = conectar(tmp_path / "cnd.db")
        criar_schema(conn)

        _lote, leitura = importar(conn, carteira, "GO", ["GO"], ORGAO)

        assert len(leitura.itens) == 4
        assert leitura.rejeitados == []
        assert {i.orgao for i in leitura.itens} == {ORGAO}
        assert {i.tipo_documento for i in leitura.itens} == {"CNPJ"}

    def test_o_orgao_esta_ligado_no_config(self):
        """Sem `ativo = true` a automação aparece na tela como indisponível
        e o orquestrador não pega os itens dela."""
        cfg = carregar()

        assert ORGAO in cfg.orgaos, "faltou a seção [orgaos.SEFAZ_GO]"
        assert ORGAO in {o.codigo for o in cfg.ativos()}

    def test_o_painel_oferece_a_automacao(self):
        from cnd.web.app import automacoes

        oferecidas = {a["codigo"]: a for a in automacoes()}

        assert ORGAO in oferecidas
        assert oferecidas[ORGAO]["disponivel"], oferecidas[ORGAO]["motivo"]
        assert oferecidas[ORGAO]["aba"] == "GO"


class TestExcelNoMesmoPadrao:
    def test_as_quatro_abas_de_sempre(self, processado, tmp_path):
        from openpyxl import load_workbook

        destino = gerar(processado, Recorte(mes=tempo.agora_iso()[:7], orgao=ORGAO),
                        tmp_path / "rel.xlsx")
        livro = load_workbook(destino)

        assert livro.sheetnames == ["Resumo", "Certidões", "Pendências", "Erros"]

    def test_negativa_e_cpen_contam_como_certidao_em_maos(self, processado,
                                                          tmp_path):
        from openpyxl import load_workbook

        destino = gerar(processado, Recorte(mes=tempo.agora_iso()[:7], orgao=ORGAO),
                        tmp_path / "rel.xlsx")
        livro = load_workbook(destino)

        # 2 negativas + 1 CPEN, sem contar o cabeçalho.
        assert livro["Certidões"].max_row - 1 == 3
        # A positiva não é entregável: vai para quem precisa de providência.
        assert livro["Pendências"].max_row - 1 == 1

    def test_a_positiva_nao_aparece_como_certidao(self, processado, tmp_path):
        from openpyxl import load_workbook

        destino = gerar(processado, Recorte(mes=tempo.agora_iso()[:7], orgao=ORGAO),
                        tmp_path / "rel.xlsx")
        certidoes = load_workbook(destino)["Certidões"]

        situacoes = [linha[3] for linha in
                     certidoes.iter_rows(min_row=2, values_only=True)]
        assert "Positiva (com pendência)" not in situacoes


class TestPacoteNoMesmoPadrao:
    def test_uma_pasta_por_orgao_com_o_nome_de_gente(self, processado):
        """`SEFAZ_GO` serve ao banco; quem recebe o pacote lê `SEFAZ GOIAS`."""
        pacote = zipar_pdfs(processado, mes=tempo.agora_iso()[:7], orgao=ORGAO)

        with zipfile.ZipFile(io.BytesIO(pacote)) as arquivo:
            nomes = arquivo.namelist()

        assert all(n.startswith(f"{nome_do_orgao(ORGAO)}/") or n == "indice.csv"
                   for n in nomes), nomes
        assert nome_do_orgao(ORGAO) == "SEFAZ GOIAS"

    def test_negativa_e_cpen_em_pastas_separadas(self, processado):
        pacote = zipar_pdfs(processado, mes=tempo.agora_iso()[:7], orgao=ORGAO)

        with zipfile.ZipFile(io.BytesIO(pacote)) as arquivo:
            nomes = arquivo.namelist()

        assert sum("CERTIDOES NEGATIVAS/" in n for n in nomes) == 2
        assert sum("POSITIVAS COM EFEITO DE NEGATIVA/" in n for n in nomes) == 1
        assert "indice.csv" in nomes

    def test_a_positiva_nao_entra_no_pacote(self, processado):
        """Entregar uma positiva junto com as negativas é o erro que ninguém
        percebe até o cliente perceber."""
        pacote = zipar_pdfs(processado, mes=tempo.agora_iso()[:7], orgao=ORGAO)

        with zipfile.ZipFile(io.BytesIO(pacote)) as arquivo:
            pdfs = [n for n in arquivo.namelist() if n.endswith(".pdf")]

        assert len(pdfs) == 3, "só as 2 negativas e a CPEN"
