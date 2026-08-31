from openpyxl import load_workbook

from cnd.core import tempo
from cnd.core.modelos import Desfecho, Status
from cnd.web.relatorio import Recorte, gerar, mes_corrente
from tests.conftest import criar_job


def test_planilha_tem_quatro_abas(conn, lote, tmp_path):
    """Uma aba por pergunta da operação, e nada além.

    Eram nove — quatro delas com colunas idênticas, mudando só o desfecho,
    mais a Auditoria (uma linha por tentativa do robô), que é diagnóstico
    técnico e não relatório.
    """
    caminho = gerar(conn, Recorte(mes_corrente()), tmp_path / "abas.xlsx")

    livro = load_workbook(caminho)
    assert livro.sheetnames == ["Resumo", "Certidões", "Pendências", "Erros"]


def test_desfecho_vira_coluna_com_nome_legivel(conn, lote, tmp_path):
    """Com a aba juntando desfechos, o desfecho tem de aparecer na linha —
    e escrito em português, não como 'PENDENCIA_MANUAL'."""
    for documento, desfecho, nome in (
        ("11222333000181", Desfecho.CPEN, "EMPRESA PARCELADA"),
        ("11444777000161", Desfecho.PENDENCIA_MANUAL, "EMPRESA SEM DADOS"),
    ):
        job = criar_job(conn, lote, documento=documento, orgao="RFB_PJ",
                        nome=nome)
        conn.execute("UPDATE job SET status = ?, desfecho = ? WHERE id = ?",
                     (Status.DONE, desfecho, job))

    livro = load_workbook(
        gerar(conn, Recorte(mes_corrente()), tmp_path / "colunas.xlsx"))

    certidoes = livro["Certidões"]
    assert certidoes["A1"].value == "Empresa" and certidoes["D1"].value == "Tipo"
    assert certidoes["A2"].value == "EMPRESA PARCELADA"
    assert certidoes["D2"].value == "Positiva c/ efeito de negativa"

    pendencias = livro["Pendências"]
    assert pendencias["D1"].value == "Situação"
    assert pendencias["A2"].value == "EMPRESA SEM DADOS"
    assert pendencias["D2"].value == "Informações insuficientes"


def test_positiva_aparece_em_pendencias_com_mensagem_corrigida(conn, lote,
                                                               tmp_path):
    job = criar_job(conn, lote, documento="40539572000168",
                    orgao="SEFAZ_GO", nome="LAMON CONSTRUTORA")
    conn.execute(
        "UPDATE job SET status = ?, desfecho = ?, atualizado_em = ? WHERE id = ?",
        (Status.DONE, Desfecho.POSITIVA, "2026-08-31T18:57:31.858Z", job),
    )
    conn.execute(
        """
        INSERT INTO tentativa
            (job_id, numero, iniciada_em, finalizada_em, desfecho,
             mensagem_portal, worker)
        VALUES (?, 1, ?, ?, ?, ?, 0)
        """,
        (
            job,
            "2026-08-31T18:57:30.738Z",
            "2026-08-31T18:57:31.856Z",
            Desfecho.POSITIVA,
            "SEFAZ-GO emitiu PDF apos confirmar nome do contribuinte",
        ),
    )

    pendencias = load_workbook(
        gerar(conn, Recorte("2026-08"), tmp_path / "positiva.xlsx")
    )["Pendências"]

    assert pendencias["A2"].value == "LAMON CONSTRUTORA"
    assert pendencias["B2"].value == "40.539.572/0001-68"
    assert pendencias["D2"].value == "Positiva (com pendência)"
    assert pendencias["F2"].value == (
        "SEFAZ-GO emitiu PDF após confirmar o nome do contribuinte."
    )


def test_aproveitada_entra_com_a_certidao_que_reusou(conn, lote, tmp_path):
    """A APROVEITADA vale como certidão em mãos e traz emissão, validade,
    código e arquivo — a do job reenfileirado, que não reemitiu.

    Ela ficava numa aba própria, no formato das pendências: sem nenhuma
    dessas colunas. Quem somava o que tinha para entregar deixava de
    contá-la, embora o PDF já estivesse no pacote."""
    pdf = tmp_path / "certidao.pdf"
    pdf.write_bytes(b"%PDF")
    job = criar_job(conn, lote, documento="11222333000181",
                    orgao="RFB_PJ", nome="EMPRESA REPETIDA")
    conn.execute(
        """
        INSERT INTO certidao
            (job_id, tipo, emitida_em, valida_ate, codigo_controle,
             caminho_pdf, sha256)
        VALUES (?, ?, ?, '2027-02-03', 'ABC123', ?, 'x')
        """,
        (job, Desfecho.NEGATIVA, tempo.agora_iso(), str(pdf)),
    )
    conn.execute("UPDATE job SET status = ?, desfecho = ? WHERE id = ?",
                 (Status.DONE, Desfecho.APROVEITADA, job))

    livro = load_workbook(
        gerar(conn, Recorte(mes_corrente()), tmp_path / "aproveitada.xlsx"))

    certidoes = livro["Certidões"]
    assert certidoes["D2"].value == "Já emitida no mês"
    assert certidoes["F2"].value == "03/02/2027"
    assert certidoes["G2"].value == "ABC123"
    assert certidoes["H2"].value == "certidao.pdf"


def test_recorte_por_orgao_deixa_os_outros_de_fora(conn, lote, tmp_path):
    """Quem filtra "Receita Federal" e pede a planilha espera receber a
    Receita Federal — antes vinha o lote inteiro, sem aviso."""
    for documento, orgao, nome in (
        ("11222333000181", "RFB_PJ", "EMPRESA FEDERAL"),
        ("11444777000161", "SEFAZ_GO", "EMPRESA ESTADUAL"),
    ):
        job = criar_job(conn, lote, documento=documento, orgao=orgao,
                        nome=nome)
        conn.execute("UPDATE job SET status = ?, desfecho = ? WHERE id = ?",
                     (Status.DONE, Desfecho.NEGATIVA, job))

    caminho = gerar(conn, Recorte(mes_corrente(), "RFB_PJ"),
                    tmp_path / "federal.xlsx")

    certidoes = load_workbook(caminho)["Certidões"]
    nomes = [linha[0] for linha in certidoes.iter_rows(min_row=2,
                                                       values_only=True)]
    assert "EMPRESA FEDERAL" in nomes
    assert "EMPRESA ESTADUAL" not in nomes


def test_recorte_aceita_varias_automacoes(conn, lote, tmp_path):
    for documento, orgao, nome in (
        ("11222333000181", "RFB_PJ", "EMPRESA FEDERAL"),
        ("11444777000161", "SEFAZ_GO", "EMPRESA ESTADUAL"),
        ("11888777000165", "CRF", "EMPRESA FGTS"),
    ):
        job = criar_job(conn, lote, documento=documento, orgao=orgao,
                        nome=nome)
        conn.execute(
            "UPDATE job SET status = ?, desfecho = ?, atualizado_em = ? "
            "WHERE id = ?",
            (Status.DONE, Desfecho.NEGATIVA, "2026-08-20T10:00:00.000Z", job),
        )

    livro = load_workbook(
        gerar(conn, Recorte("2026-08", ("RFB_PJ", "SEFAZ_GO")),
              tmp_path / "selecionadas.xlsx")
    )

    nomes = [linha[0] for linha in livro["Certidões"].iter_rows(
        min_row=2, values_only=True)]
    assert set(nomes) == {"EMPRESA FEDERAL", "EMPRESA ESTADUAL"}
    assert "EMPRESA FGTS" not in nomes
    assert livro["Resumo"]["B3"].value == "RFB_PJ, SEFAZ_GO"


def test_arquivo_pdf_nao_vira_link_quebrado(conn, lote, tmp_path):
    job = criar_job(conn, lote, documento="11222333000181",
                    orgao="RFB_PJ", nome="EMPRESA COM PDF")
    pdf = tmp_path / "certidao.pdf"
    pdf.write_bytes(b"%PDF")
    conn.execute(
        "UPDATE job SET status = ?, desfecho = ? WHERE id = ?",
        (Status.DONE, Desfecho.NEGATIVA, job),
    )
    conn.execute(
        """
        INSERT INTO certidao
            (job_id, tipo, emitida_em, valida_ate, codigo_controle,
             caminho_pdf, sha256)
        VALUES (?, ?, ?, '2027-02-03', 'ABC123', ?, 'x')
        """,
        (job, Desfecho.NEGATIVA, tempo.agora_iso(), str(pdf)),
    )

    caminho = gerar(conn, Recorte(mes_corrente()), tmp_path / "pdf.xlsx")

    certidoes = load_workbook(caminho)["Certidões"]
    assert certidoes["H2"].value == "certidao.pdf"
    assert certidoes["H2"].hyperlink is None


def test_sem_orgao_traz_todos(conn, lote, tmp_path):
    for documento, orgao in (("11222333000181", "RFB_PJ"),
                             ("11444777000161", "SEFAZ_GO")):
        job = criar_job(conn, lote, documento=documento, orgao=orgao,
                        nome=f"EMPRESA {orgao}")
        conn.execute("UPDATE job SET status = ?, desfecho = ? WHERE id = ?",
                     (Status.DONE, Desfecho.NEGATIVA, job))

    caminho = gerar(conn, Recorte(mes_corrente()), tmp_path / "tudo.xlsx")

    certidoes = load_workbook(caminho)["Certidões"]
    assert certidoes.max_row == 3, "cabeçalho + duas empresas"


def test_o_resumo_conta_o_mesmo_que_as_abas(conn, lote, tmp_path):
    """O Resumo pedia os números SEM recorte nenhum: em agosto, a linha
    "Total" somava julho junto e não batia com a aba logo abaixo.

    Um resumo que discorda das abas do próprio arquivo é o pior defeito
    possível num relatório — quem confere para de confiar no arquivo
    inteiro, inclusive na parte que estava certa.
    """
    for documento, quando, nome in (
        ("11222333000181", "2026-07-20T10:00:00.000Z", "EMPRESA DE JULHO"),
        ("11444777000161", "2026-08-20T10:00:00.000Z", "EMPRESA DE AGOSTO"),
    ):
        job = criar_job(conn, lote, documento=documento, orgao="RFB_PJ",
                        nome=nome)
        conn.execute(
            "UPDATE job SET status = ?, desfecho = ?, atualizado_em = ? "
            "WHERE id = ?",
            (Status.DONE, Desfecho.NEGATIVA, quando, job))

    livro = load_workbook(
        gerar(conn, Recorte("2026-08"), tmp_path / "agosto.xlsx"))

    # Linha 6 é o cabeçalho da tabela por órgão; a 7 é o primeiro órgão.
    resumo = livro["Resumo"]
    assert resumo["A6"].value == "Órgão" and resumo["B6"].value == "Total"
    assert resumo["A7"].value == "RFB_PJ"
    assert resumo["B7"].value == 1, "total do MÊS, não do banco inteiro"
    assert resumo["C7"].value == 1
    assert livro["Certidões"].max_row == 2, "cabeçalho + a empresa de agosto"


def test_o_cabecalho_pintado_e_o_cabecalho(conn, lote, tmp_path):
    """O número fixo apontava para a primeira linha de ÓRGÃO: a Receita
    Federal saía pintada de azul escuro, como se fosse título, e o
    cabeçalho de verdade sem destaque nenhum."""
    job = criar_job(conn, lote, documento="11222333000181", orgao="RFB_PJ")
    conn.execute(
        "UPDATE job SET status = ?, desfecho = ?, atualizado_em = ? WHERE id = ?",
        (Status.DONE, Desfecho.NEGATIVA, "2026-08-20T10:00:00.000Z", job))

    resumo = load_workbook(
        gerar(conn, Recorte("2026-08"), tmp_path / "pintura.xlsx"))["Resumo"]

    assert resumo["A6"].font.color.rgb.endswith("FFFFFF"), "o cabeçalho é claro"
    assert not resumo["A7"].font.color, "a linha do órgão é linha de dados"


def test_o_resumo_nao_lista_orgao_que_nao_trabalhou_no_mes(conn, lote,
                                                           tmp_path):
    """Sem recorte, a lista de órgãos vinha do banco inteiro — e a planilha
    de agosto trazia uma linha da SEFAZ que só rodou em julho, com os
    números daquele mês."""
    job = criar_job(conn, lote, documento="11222333000181", orgao="SEFAZ_GO",
                    nome="EMPRESA ESTADUAL")
    conn.execute(
        "UPDATE job SET status = ?, desfecho = ?, atualizado_em = ? WHERE id = ?",
        (Status.DONE, Desfecho.NEGATIVA, "2026-07-20T10:00:00.000Z", job))

    livro = load_workbook(
        gerar(conn, Recorte("2026-08"), tmp_path / "vazio.xlsx"))

    orgaos = [linha[0] for linha in livro["Resumo"].iter_rows(
        min_row=8, max_col=1, values_only=True)]
    assert "SEFAZ_GO" not in orgaos


def test_mes_anterior_fica_de_fora(conn, lote, tmp_path):
    """A emissão é mensal: a planilha de agosto não leva o que saiu em
    julho, senão o cliente recebe certidão vencida na conferência."""
    job = criar_job(conn, lote, documento="11222333000181", orgao="RFB_PJ",
                    nome="EMPRESA ANTIGA")
    conn.execute(
        "UPDATE job SET status = ?, desfecho = ?, atualizado_em = ? WHERE id = ?",
        (Status.DONE, Desfecho.NEGATIVA, "2020-01-15T10:00:00.000Z", job))

    caminho = gerar(conn, Recorte(mes_corrente()), tmp_path / "mes.xlsx")

    certidoes = load_workbook(caminho)["Certidões"]
    assert certidoes.max_row == 1, "só o cabeçalho"
