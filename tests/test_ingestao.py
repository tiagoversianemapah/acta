"""Testes da ingestão da planilha.

Constrói um Excel na hora, com casos reais que a base da Mapah tem:
documento válido, dígito errado, CPF na aba de PJ, duplicata e linha vazia.
"""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from cnd.ingestao.planilha import importar, ler

VALIDO_A = "11.222.333/0001-81"
VALIDO_B = "12.ABC.345/01DE-35"     # formato alfanumérico novo
DV_ERRADO = "11.222.333/0001-82"
CPF = "529.982.247-25"              # pessoa física na aba de PJ


@pytest.fixture
def planilha(tmp_path):
    livro = Workbook()
    livro.remove(livro.active)

    aba = livro.create_sheet("RFB ")      # com espaço, como na planilha real
    aba.append(["Empresa", "CNPJ"])
    aba.append(["EMPRESA A", VALIDO_A])
    aba.append(["EMPRESA B", VALIDO_B])
    aba.append(["EMPRESA C", DV_ERRADO])
    aba.append(["PESSOA FISICA", CPF])
    aba.append(["EMPRESA A DE NOVO", VALIDO_A])
    aba.append([None, None])
    aba.append(["EMPRESA D", 11222333000181])   # número, sem máscara

    outra = livro.create_sheet("CRF")
    outra.append(["Empresa", "CNPJ"])
    outra.append(["EMPRESA A", VALIDO_A])

    caminho = tmp_path / "teste.xlsx"
    livro.save(caminho)
    return caminho


def test_le_apenas_a_aba_pedida(planilha):
    leitura = ler(planilha, ["RFB"])

    assert {i.orgao for i in leitura.itens} == {"RFB_PJ"}


def test_aceita_validos_em_qualquer_formato(planilha):
    leitura = ler(planilha, ["RFB"])
    documentos = {i.documento for i in leitura.itens}

    assert "11222333000181" in documentos
    assert "12ABC34501DE35" in documentos, "o CNPJ alfanumérico precisa passar"


def test_rejeita_dv_errado_com_motivo(planilha):
    leitura = ler(planilha, ["RFB"])

    motivos = [r.motivo for r in leitura.rejeitados if r.valor_original == DV_ERRADO]
    assert motivos and "dígito verificador" in motivos[0]


def test_rejeita_cpf_na_aba_de_pj(planilha):
    leitura = ler(planilha, ["RFB"])

    assert any(r.valor_original == CPF for r in leitura.rejeitados)


def test_duplicata_entra_uma_vez_so(planilha):
    leitura = ler(planilha, ["RFB"])

    assert sum(1 for i in leitura.itens if i.documento == "11222333000181") == 1
    assert any("duplicado" in r.motivo for r in leitura.rejeitados)


def test_linha_vazia_e_ignorada(planilha):
    leitura = ler(planilha, ["RFB"])

    assert not any(r.valor_original == "None" for r in leitura.rejeitados)


def test_rejeitado_aponta_a_linha_da_planilha(planilha):
    leitura = ler(planilha, ["RFB"])

    rejeitado = next(r for r in leitura.rejeitados if r.valor_original == DV_ERRADO)
    assert rejeitado.linha == 4, "quem for corrigir precisa achar a célula"
    assert rejeitado.aba == "RFB"


def test_importar_cria_lote_empresas_e_jobs(conn, planilha):
    lote_id, leitura = importar(conn, planilha, "teste", ["RFB"])

    jobs = conn.execute("SELECT COUNT(*) AS n FROM job WHERE lote_id = ?",
                        (lote_id,)).fetchone()["n"]
    assert jobs == len(leitura.itens)
    assert conn.execute("SELECT COUNT(*) AS n FROM empresa").fetchone()["n"] == len(leitura.itens)


def test_todo_job_nasce_pendente(conn, planilha):
    lote_id, _ = importar(conn, planilha, "teste", ["RFB"])

    situacoes = {linha["status"] for linha in
                 conn.execute("SELECT status FROM job WHERE lote_id = ?", (lote_id,))}
    assert situacoes == {"PENDING"}


def test_reimportar_nao_duplica_empresa(conn, planilha):
    importar(conn, planilha, "primeira", ["RFB"])
    importar(conn, planilha, "segunda", ["RFB"])

    empresas = conn.execute("SELECT COUNT(*) AS n FROM empresa").fetchone()["n"]
    lotes = conn.execute("SELECT COUNT(*) AS n FROM lote").fetchone()["n"]
    assert lotes == 2
    # A planilha tem 2 documentos únicos válidos (EMPRESA D repete o CNPJ da A).
    assert empresas == 2, "a mesma empresa não pode virar cadastro novo a cada lote"


def test_duas_abas_geram_orgaos_diferentes(conn, planilha):
    lote_id, _ = importar(conn, planilha, "teste", ["RFB", "CRF"])

    orgaos = {linha["orgao"] for linha in
              conn.execute("SELECT DISTINCT orgao FROM job WHERE lote_id = ?", (lote_id,))}
    assert orgaos == {"RFB_PJ", "CRF"}


def test_importar_aposenta_os_envios_antigos(conn, planilha, tmp_path):
    """Planilha nova entra, a sexta mais antiga sai — com os PDFs dela.

    A limpeza acontece na importação porque é o único momento em que a
    máquina ganha um envio: sem isso o banco crescia desde a primeira
    rodada, e o disco junto (22/09/2026). Só sai envio TERMINADO: por isso
    os itens são fechados aqui, como o robô faria.
    """
    from cnd.core.modelos import Status

    certidoes = tmp_path / "certidoes"
    certidoes.mkdir()
    primeiros = []
    for n in range(6):
        lote, _leitura = importar(conn, planilha, f"envio {n}",
                                  pasta_certidoes=certidoes)
        conn.execute("UPDATE job SET status = ? WHERE lote_id = ?",
                     (Status.DONE, lote))
        primeiros.append(lote)

    importar(conn, planilha, "envio 7", pasta_certidoes=certidoes)
    vivos = [linha["id"] for linha in conn.execute("SELECT id FROM lote")]

    # Sobram os cinco terminados mais recentes e o que acabou de entrar,
    # que ainda tem fila: só o primeiro envio saiu.
    assert primeiros[0] not in vivos, "o mais antigo saiu"
    assert vivos == primeiros[1:] + [vivos[-1]]


def test_importar_nao_apaga_planilha_com_fila_aberta(conn, planilha, tmp_path):
    """Importar não pode levar junto trabalho que ninguém mandou parar."""
    from cnd.core.modelos import Status

    certidoes = tmp_path / "certidoes"
    certidoes.mkdir()
    antigo, _leitura = importar(conn, planilha, "com fila",
                                pasta_certidoes=certidoes)
    for n in range(6):
        lote, _l = importar(conn, planilha, f"envio {n}",
                            pasta_certidoes=certidoes)
        conn.execute("UPDATE job SET status = ? WHERE lote_id = ?",
                     (Status.DONE, lote))
    importar(conn, planilha, "o que dispara a limpeza",
             pasta_certidoes=certidoes)

    vivos = [linha["id"] for linha in conn.execute("SELECT id FROM lote")]

    assert antigo in vivos, "envio com item pendente fica, por mais velho"
