"""Aba e automação são escolhas separadas.

A aba diz ONDE estão os dados; a automação diz O QUE fazer com eles. Antes
a aba decidia as duas coisas, e uma planilha de cliente com a aba chamada
"Clientes GO" era simplesmente ignorada — o envio dizia "0 itens entraram
na fila" sem explicar por quê.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from cnd.ingestao.planilha import ORGAO_PARA_TIPO, abas_da_planilha, ler

CNPJ_A = "33949051000113"
CNPJ_B = "06031097000186"
CPF_OK = "11144477735"


@pytest.fixture
def planilha(tmp_path) -> Path:
    livro = Workbook()
    rfb = livro.active
    rfb.title = "RFB"
    rfb.append(["Empresa", "Documento"])
    rfb.append(["EMPRESA A", CNPJ_A])

    crf = livro.create_sheet("CRF")
    crf.append(["Empresa", "Documento"])
    crf.append(["EMPRESA B", CNPJ_B])

    livre = livro.create_sheet("Clientes GO")
    livre.append(["Empresa", "Documento"])
    livre.append(["EMPRESA C", CNPJ_A])

    livro.create_sheet("Instruções").append(["leia-me"])
    caminho = tmp_path / "carteira.xlsx"
    livro.save(caminho)
    return caminho


def test_lista_as_abas_com_a_contagem(planilha):
    """A tela precisa mostrar o que o arquivo tem ANTES de importar."""
    achado = dict(abas_da_planilha(planilha))
    assert achado["RFB"] == 1
    assert achado["CRF"] == 1
    assert achado["Clientes GO"] == 1
    assert "Instruções" in achado, "aba sem automação também aparece na lista"


def test_aba_com_nome_livre_roda_a_automacao_escolhida(planilha):
    """O caso que não funcionava: nome de aba que não é código de órgão."""
    leitura = ler(planilha, abas=["Clientes GO"], orgao="CRF")
    assert [i.orgao for i in leitura.itens] == ["CRF"]
    assert [i.documento for i in leitura.itens] == [CNPJ_A]


def test_a_automacao_manda_sobre_o_nome_da_aba(planilha):
    """Aba chamada RFB pode rodar FGTS, se for o que se pediu."""
    leitura = ler(planilha, abas=["RFB"], orgao="CRF")
    assert [i.orgao for i in leitura.itens] == ["CRF"]


def test_sem_automacao_vale_o_atalho_antigo(planilha):
    """A linha de comando e a planilha com abas nomeadas continuam iguais."""
    leitura = ler(planilha)
    assert {i.orgao for i in leitura.itens} == {"RFB_PJ", "CRF"}
    assert not any(i.orgao == "Clientes GO" for i in leitura.itens)


def test_o_tipo_do_documento_vem_do_orgao(planilha, tmp_path):
    """Com aba livre não há como deduzir CNPJ ou CPF do nome dela.

    Quem sabe é a automação: o FGTS consulta empregador (CNPJ), a Receita
    PF consulta CPF. Sem isso, o CPF válido abaixo seria recusado por não
    ser um CNPJ.
    """
    livro = Workbook()
    aba = livro.active
    aba.title = "Pessoas"
    aba.append(["Nome", "Documento", "Data de Nascimento"])
    aba.append(["FULANO", CPF_OK, "22/02/1948"])
    caminho = tmp_path / "pf.xlsx"
    livro.save(caminho)

    assert ORGAO_PARA_TIPO["RFB_PF"] == "CPF"
    leitura = ler(caminho, abas=["Pessoas"], orgao="RFB_PF")
    assert [i.documento for i in leitura.itens] == [CPF_OK]
    assert not leitura.rejeitados
