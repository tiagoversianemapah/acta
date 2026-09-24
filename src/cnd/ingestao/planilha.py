"""Lê a planilha Excel e cria os jobs no banco.

Esta é a única parte do sistema que conhece Excel. Se um dia a fonte
virar uma API ou um sistema contábil, troca-se só este arquivo.
"""
from __future__ import annotations

import re
import sqlite3
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from cnd.core.documentos import DocumentoInvalido, validar
from cnd.infra import limpeza
from cnd.infra.log import obter

log = obter("ingestao")

# aba da planilha -> (código do órgão, tipo de documento esperado)
ABA_PARA_ORGAO: dict[str, tuple[str, str]] = {
    "RFB": ("RFB_PJ", "CNPJ"),
    "CRF": ("CRF", "CNPJ"),
    "CPF": ("RFB_PF", "CPF"),
    "GO": ("SEFAZ_GO", "CNPJ"),
    "MA": ("SEFAZ_MA", "CNPJ"),
    "MT": ("SEFAZ_MT", "CNPJ"),
    "GOIANIA": ("GOIANIA", "CNPJ"),
    "VITORIA": ("VITORIA", "CNPJ"),
    "SAO LUIS": ("SAO_LUIS", "CNPJ"),
    "DF": ("SEFAZ_DF", "CNPJ"),
    "ES": ("SEFAZ_ES", "CNPJ"),
    "SP": ("SEFAZ_SP", "CNPJ"),
}


# O tipo de documento passa a vir do ÓRGÃO, e não da aba. Com aba e
# automação escolhidas em separado, a aba deixou de dizer o que ela contém:
# uma aba "Clientes GO" pode rodar FGTS. Quem sabe se valida CNPJ ou CPF é a
# automação — o FGTS consulta empregador (CNPJ), a Receita PF consulta CPF.
ORGAO_PARA_TIPO: dict[str, str] = {
    orgao: tipo for orgao, tipo in ABA_PARA_ORGAO.values()
}

# Automações cujo formulário não emite sem a data de nascimento. A linha que
# chega sem ela é recusada JÁ na importação: deixá-la entrar só adiaria o
# problema para o robô, que abriria o portal para não ter o que digitar.
EXIGEM_NASCIMENTO = frozenset({"RFB_PF"})

RE_DATA_BR = re.compile(r"^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$")
RE_DATA_ISO = re.compile(r"^\s*(\d{4})-(\d{2})-(\d{2})")


def abas_da_planilha(caminho: Path) -> list[tuple[str, int]]:
    """Nome e quantidade de linhas de cada aba, para a tela oferecer a escolha.

    A tela precisa disto ANTES de importar: as abas de um arquivo só existem
    depois de abri-lo, e quem envia precisa ver o que tem lá para escolher.
    `read_only` porque planilha de carteira inteira não cabe confortavelmente
    na memória de uma máquina de escritório.
    """
    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        return [(nome, max((livro[nome].max_row or 1) - 1, 0))
                for nome in livro.sheetnames]
    finally:
        livro.close()


def _sem_acento(texto) -> str:
    ascii_puro = (unicodedata.normalize("NFKD", str(texto or ""))
                  .encode("ascii", "ignore").decode("ascii"))
    return " ".join(ascii_puro.lower().split())


def _ler_data(valor) -> date | None:
    """Data de uma célula: a do Excel ou texto dd/mm/aaaa. Senão, None."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "")
    try:
        if achado := RE_DATA_BR.match(texto):
            dia, mes, ano = (int(g) for g in achado.groups())
            return date(ano, mes, dia)
        if achado := RE_DATA_ISO.match(texto):
            ano, mes, dia = (int(g) for g in achado.groups())
            return date(ano, mes, dia)
    except ValueError:
        return None
    return None


def _coluna_nascimento(cabecalho) -> int | None:
    """A coluna cujo título fala em nascimento.

    Só pelo título, e de propósito: adivinhar pela primeira data da linha
    pegaria a competência que as abas da carteira trazem na coluna C — e
    numa planilha reimportada meses depois ela já parece data de nascimento.
    Substring, e não palavra inteira, para aceitar o que vem de sistema:
    "DT_NASCIMENTO", "DataNascimento", "Nasc.". E "D.N.", que é como a
    planilha de pessoas físicas do escritório chama a coluna (17/09/2026).
    """
    for indice, titulo in enumerate(cabecalho or ()):
        normalizado = _sem_acento(titulo)
        if "nasc" in normalizado or re.sub(r"[^a-z]", "", normalizado) == "dn":
            return indice
    return None


def _nascimento_da_linha(linha, coluna: int | None) -> date | None:
    if coluna is None or len(linha) <= coluna:
        return None
    return _ler_data(linha[coluna])


@dataclass(frozen=True)
class Item:
    orgao: str
    tipo_documento: str
    documento: str
    nome: str
    data_nascimento: date | None = None


@dataclass(frozen=True)
class Rejeitado:
    aba: str
    linha: int
    valor_original: str
    nome: str
    motivo: str


@dataclass
class Leitura:
    itens: list[Item] = field(default_factory=list)
    rejeitados: list[Rejeitado] = field(default_factory=list)


def ler(caminho: Path, abas: list[str] | None = None,
        orgao: str | None = None) -> Leitura:
    """Lê a planilha e devolve os válidos e os rejeitados.
    Não mexe no banco: só lê e organiza.

    Com `orgao`, a AUTOMAÇÃO escolhida manda: as abas pedidas são lidas
    inteiras para esse órgão, tenham o nome que tiverem. É o que permite
    "aba Clientes GO, automação FGTS" — a aba diz onde estão os dados, a
    automação diz o que fazer com eles.

    Sem `orgao`, vale o atalho antigo: a aba chamada RFB é da Receita, a
    CRF é do FGTS. Continua servindo à linha de comando e à planilha que
    já vem com as abas nomeadas.
    """
    resultado = Leitura()
    livro = load_workbook(caminho, read_only=True, data_only=True)
    vistos: set[tuple[str, str]] = set()
    pedidas = {a.strip().upper() for a in (abas or [])}

    for nome_aba in livro.sheetnames:
        chave = nome_aba.strip().upper()
        if orgao:
            if pedidas and chave not in pedidas:
                continue
            destino, tipo = orgao, ORGAO_PARA_TIPO.get(orgao, "CNPJ")
        else:
            if chave not in ABA_PARA_ORGAO:
                continue
            if pedidas and chave not in pedidas:
                continue
            destino, tipo = ABA_PARA_ORGAO[chave]
        orgao_da_aba = destino
        linhas = livro[nome_aba].iter_rows(values_only=True)
        # linha 1 é o cabeçalho: não vira item, mas diz onde está o nascimento
        coluna_nascimento = _coluna_nascimento(next(linhas, None))

        for numero, linha in enumerate(linhas, start=2):
            if not linha or all(c is None for c in linha):
                continue

            nome = str(linha[0] or "").strip()
            bruto = linha[1] if len(linha) > 1 else None

            try:
                documento = validar(bruto, tipo)
            except DocumentoInvalido as erro:
                resultado.rejeitados.append(
                    Rejeitado(chave, numero, str(bruto), nome, str(erro))
                )
                continue

            if (orgao_da_aba, documento) in vistos:
                resultado.rejeitados.append(
                    Rejeitado(chave, numero, str(bruto), nome,
                              "duplicado na planilha (mantida a primeira ocorrência)")
                )
                continue

            nascimento = (_nascimento_da_linha(linha, coluna_nascimento)
                          if tipo == "CPF" else None)
            if orgao_da_aba in EXIGEM_NASCIMENTO:
                # Dois motivos, e não um: coluna faltando se resolve uma vez
                # na planilha inteira; célula vazia, linha a linha.
                if coluna_nascimento is None:
                    resultado.rejeitados.append(
                        Rejeitado(chave, numero, str(bruto), nome,
                                  "a aba não tem a coluna 'Data de Nascimento': "
                                  "a Receita não emite a certidão de CPF sem ela")
                    )
                    continue
                if nascimento is None:
                    resultado.rejeitados.append(
                        Rejeitado(chave, numero, str(bruto), nome,
                                  "data de nascimento vazia ou ilegível "
                                  "(use dd/mm/aaaa)")
                    )
                    continue
                if nascimento > date.today():
                    resultado.rejeitados.append(
                        Rejeitado(chave, numero, str(bruto), nome,
                                  f"data de nascimento no futuro: "
                                  f"{nascimento:%d/%m/%Y}")
                    )
                    continue

            vistos.add((orgao_da_aba, documento))
            resultado.itens.append(
                Item(orgao_da_aba, tipo, documento, nome or documento,
                     nascimento))

    livro.close()
    return resultado


def ler_pares(caminho: Path, pares: list[tuple[str, str]]) -> Leitura:
    """Várias abas, cada uma com a sua automação, numa leitura só.

    O mesmo documento em duas abas da MESMA automação entra uma vez: a
    segunda ocorrência vira recusa, como o repetido dentro de uma aba.
    """
    total = Leitura()
    vistos: set[tuple[str, str]] = set()
    for aba, orgao in pares:
        parcial = ler(caminho, [aba], orgao)
        total.rejeitados.extend(parcial.rejeitados)
        for item in parcial.itens:
            chave = (item.orgao, item.documento)
            if chave in vistos:
                total.rejeitados.append(Rejeitado(
                    aba.strip().upper(), 0, item.documento, item.nome,
                    "duplicado em outra aba da mesma automação "
                    "(mantida a primeira ocorrência)"))
                continue
            vistos.add(chave)
            total.itens.append(item)
    return total


def importar(conn: sqlite3.Connection, caminho: Path, descricao: str,
             abas: list[str] | None = None, orgao: str | None = None,
             arquivo_origem: str | None = None,
             pares: list[tuple[str, str]] | None = None,
             pasta_certidoes: Path | None = None) -> tuple[int, Leitura]:
    """Lê a planilha e grava lote + empresas + jobs no banco.

    Tudo de uma vez só: ou o lote inteiro entra, ou nada entra.

    `pares` importa várias abas, cada uma com a sua automação, num lote SÓ.
    A planilha é uma; a tela, o controle de fila e a entrega a tratam como
    uma. Importada aba por aba, a carteira virava um lote por automação com
    o mesmo nome, e o painel mostrava só um deles (17/09/2026).

    Entrando um envio novo, os que passarem dos cinco mais recentes são
    apagados com os PDFs deles — ver `limpeza.aposentar_envios`. É o mesmo
    momento em que o painel já aposentava o ARQUIVO da planilha antiga; o
    banco é que guardava tudo desde a primeira rodada.
    """
    leitura = ler_pares(caminho, pares) if pares else ler(caminho, abas, orgao)

    conn.execute("BEGIN")
    try:
        cursor = conn.execute(
            "INSERT INTO lote (descricao, arquivo_origem) VALUES (?, ?)",
            (descricao, arquivo_origem or caminho.name),
        )
        lote_id = cursor.lastrowid

        for item in leitura.itens:
            # Importação sem data não apaga a que já está guardada: o CPF que
            # entrar por outra automação não pode desarmar a Receita PF.
            conn.execute(
                "INSERT INTO empresa (documento, tipo_documento, nome, data_nascimento) "
                "VALUES (?, ?, ?, ?) "
                "ON CONFLICT (documento) DO UPDATE SET nome = excluded.nome, "
                "data_nascimento = COALESCE(excluded.data_nascimento, "
                "empresa.data_nascimento)",
                (item.documento, item.tipo_documento, item.nome,
                 item.data_nascimento.isoformat() if item.data_nascimento
                 else None),
            )
            empresa_id = conn.execute(
                "SELECT id FROM empresa WHERE documento = ?", (item.documento,)
            ).fetchone()["id"]

            conn.execute(
                "INSERT INTO job (lote_id, empresa_id, orgao) VALUES (?, ?, ?) "
                "ON CONFLICT (lote_id, empresa_id, orgao) DO NOTHING",
                (lote_id, empresa_id, item.orgao),
            )
        conn.execute("COMMIT")
    except Exception:
        conn.execute("ROLLBACK")
        raise

    # Depois do COMMIT, e fora da transação do envio: falhar ao apagar o que
    # é velho não pode derrubar a importação que acabou de dar certo.
    try:
        limpeza.aposentar_envios(conn, pasta_certidoes)
    except Exception as erro:
        log.warning("aposentadoria_falhou", extra={"erro": str(erro)[:300]})

    return lote_id, leitura


def _main() -> None:
    import argparse

    from cnd.infra.db import conectar, criar_schema

    parser = argparse.ArgumentParser(description="Importa uma planilha e cria um lote.")
    parser.add_argument("planilha", type=Path)
    parser.add_argument("--descricao", default="")
    parser.add_argument("--abas", nargs="*", default=["RFB"])
    args = parser.parse_args()

    conn = conectar()
    criar_schema(conn)

    descricao = args.descricao or f"Importação de {args.planilha.name}"
    abas = [a.upper() for a in args.abas] if args.abas else None
    lote_id, leitura = importar(conn, args.planilha, descricao, abas)

    print(f"\nLote #{lote_id} criado a partir de {args.planilha.name}")
    print(f"  jobs criados : {len(leitura.itens)}")
    print(f"  rejeitados   : {len(leitura.rejeitados)}")

    por_orgao: dict[str, int] = {}
    for item in leitura.itens:
        por_orgao[item.orgao] = por_orgao.get(item.orgao, 0) + 1
    for orgao, total in sorted(por_orgao.items()):
        print(f"    {orgao:10s} {total:5d}")

    if leitura.rejeitados:
        print("\n  Primeiros rejeitados:")
        for r in leitura.rejeitados[:10]:
            print(f"    [{r.aba} linha {r.linha}] {r.valor_original!r}: {r.motivo}")
        if len(leitura.rejeitados) > 10:
            print(f"    ... e mais {len(leitura.rejeitados) - 10}")

    conn.close()


if __name__ == "__main__":
    _main()
